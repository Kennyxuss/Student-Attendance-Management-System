import sys
try:
    from pptx import Presentation
except ImportError:
    print("python-pptx not installed. Please install it first.")
    sys.exit(1)

prs = Presentation()

# Slide 1
slide1 = prs.slides.add_slide(prs.slide_layouts[1])
slide1.shapes.title.text = "Student Attendance Management System (SAMS)"
tf1 = slide1.shapes.placeholders[1].text_frame
tf1.text = "The Problem with Paper: Sign-in sheets are easily lost, damaged, and impossible to quickly search."
tf1.add_paragraph().text = "Wasted Time: Instructors spend too much time manually calculating chronic absences."
tf1.add_paragraph().text = "No Real-Time Data: Center administrators lack instant visibility into who is actually present across different classrooms."

# Slide 2
slide2 = prs.slides.add_slide(prs.slide_layouts[1])
slide2.shapes.title.text = "A Centralized Digital Solution"
tf2 = slide2.shapes.placeholders[1].text_frame
tf2.text = "Role-Based Access: Distinct management workflows for Center Admins and Class Instructors."
tf2.add_paragraph().text = "Core Entities: Seamlessly connects Students, Instructors, Classes, and daily Attendance Logs."
tf2.add_paragraph().text = "Fast & Responsive: Built as a modern, single-page application where data updates instantly without refreshing the page."

# Slide 3
slide3 = prs.slides.add_slide(prs.slide_layouts[1])
slide3.shapes.title.text = "How It Works Under the Hood"
tf3 = slide3.shapes.placeholders[1].text_frame
tf3.text = "Frontend: HTML5, CSS3, Bootstrap 5, and Vanilla JavaScript (fetch API)."
tf3.add_paragraph().text = "Backend: Custom Python HTTP Server (Built from scratch without heavy frameworks)."
tf3.add_paragraph().text = "Strict Security: Boundary validation (guards.py) intercepts blank or malicious data and safely rejects it with a 422 error."
tf3.add_paragraph().text = "Data Storage: Scalable JSON repository, fully decoupled and prepared for MySQL migration."

# Slide 4
slide4 = prs.slides.add_slide(prs.slide_layouts[1])
slide4.shapes.title.text = "SAMS in Action"
tf4 = slide4.shapes.placeholders[1].text_frame
tf4.text = "Happy Path: Adding a new student and taking batch Roll Call."
tf4.add_paragraph().text = "Graceful Failure: Demonstrating how the system catches blank inputs safely."
tf4.add_paragraph().text = "User Experience (UX): Global loading spinners prevent duplicate submissions."
tf4.add_paragraph().text = "Interactive Feedback: Real-time 'Toast' notifications inform the user without crashing the app."

# Slide 5
slide5 = prs.slides.add_slide(prs.slide_layouts[1])
slide5.shapes.title.text = "Lessons Learned & Quality Assurance"
tf5 = slide5.shapes.placeholders[1].text_frame
tf5.text = "Organized Workflow: All tasks and bugs were triaged using a strict GitHub Project Board."
tf5.add_paragraph().text = "Adversarial Testing: We intentionally tried to break the app in Week 10 to find edge cases."
tf5.add_paragraph().text = "Security Patched: We identified a Cross-Site Scripting (XSS) vulnerability in our tables and successfully patched it."
tf5.add_paragraph().text = "Final Takeaway: AI is an incredible tool for accelerating development, but human review is absolutely necessary for true security and edge-case handling."

# Save the presentation
prs.save('docs/SAMS_Final_Presentation.pptx')
print("Presentation generated successfully!")
